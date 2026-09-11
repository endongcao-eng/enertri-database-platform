#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
wiki_to_enertri_excel_v2.py

从英文 Wikipedia 分类页抓取能源、光伏、储能、碳中和相关术语，
通过 Wikidata 获取中/英/俄三语名称与描述，导出为 EnerTri 平台可导入的 Excel。

相比旧版新增：
1. 支持保留“缺少中文术语”的 SKIP 条目。
2. 支持把 SKIP 条目放入第二个工作表“待人工翻译_SKIP”。
3. 支持用“待译：英文术语”作为中文术语，生成可导入行。
4. 优先读取 zh-hans / zh-cn，减少繁体中文标签。

安装依赖：
    pip install requests openpyxl

示例：
    python wiki_to_enertri_excel_v2.py --target 150 --category "Category:Emissions reduction" --output EnerTri_碳中和术语_150条.xlsx --missing-zh-mode separate
    python wiki_to_enertri_excel_v2.py --target 150 --category "Category:Emissions reduction" --output EnerTri_碳中和术语_150条.xlsx --missing-zh-mode importable
    python wiki_to_enertri_excel_v2.py --target 150 --category "Category:Emissions reduction" --output EnerTri_碳中和术语_150条.xlsx --missing-zh-mode both
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from dataclasses import dataclass
from typing import Dict, Iterable, List, Optional, Set, Tuple
from urllib.parse import quote

import openpyxl
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_CATEGORIES = [
    "Category:Photovoltaics",
    "Category:Solar energy",
    "Category:Renewable energy",
    "Category:Energy storage",
    "Category:Climate change mitigation",
    "Category:Carbon dioxide",
    "Category:Greenhouse gases",
    "Category:Carbon finance",
    "Category:Emissions reduction",
]

COLUMNS = [
    "中文术语",
    "英文术语",
    "俄文术语（可选）",
    "学术版术语解释",
    "通俗版术语解释",
    "简要英文解释",
    "关键词（3–5个，用逗号分隔）",
    "相关术语",
    "应用场景",
    "公式 / 模型（如有）",
    "有无视频形式的示例或工程案例",
    "参考文献（1–2篇）",
    "填写本词条的学生姓名",
    "填写本词条的学生研究方向",
    "审核状态",
]

SKIPPED_COLUMNS = [
    "处理建议",
    "跳过原因",
    "英文术语",
    "建议中文术语",
    "俄文术语（如有）",
    "英文摘要 / 描述",
    "来源分类",
    "Wikidata QID",
    "英文 Wikipedia 链接",
    "Wikidata 链接",
]


@dataclass
class TermRow:
    zh_term: str
    en_term: str
    ru_term: str
    academic_explanation: str
    plain_explanation: str
    short_en_explanation: str
    keywords: str
    related_terms: str
    application_scenario: str
    formula_or_model: str
    has_video_case: str
    references: str
    contributor_name: str
    contributor_research_area: str
    status: str

    def to_list(self) -> List[str]:
        return [
            self.zh_term,
            self.en_term,
            self.ru_term,
            self.academic_explanation,
            self.plain_explanation,
            self.short_en_explanation,
            self.keywords,
            self.related_terms,
            self.application_scenario,
            self.formula_or_model,
            self.has_video_case,
            self.references,
            self.contributor_name,
            self.contributor_research_area,
            self.status,
        ]


@dataclass
class SkippedRow:
    action: str
    reason: str
    en_term: str
    suggested_zh_term: str
    ru_term: str
    english_summary: str
    source_category: str
    qid: str
    enwiki_url: str
    wikidata_url: str

    def to_list(self) -> List[str]:
        return [
            self.action,
            self.reason,
            self.en_term,
            self.suggested_zh_term,
            self.ru_term,
            self.english_summary,
            self.source_category,
            self.qid,
            self.enwiki_url,
            self.wikidata_url,
        ]


class WikiEnerTriBuilder:
    def __init__(self, user_agent: str, delay: float = 0.5, timeout: int = 30, verbose: bool = True) -> None:
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": user_agent})
        self.delay = delay
        self.timeout = timeout
        self.verbose = verbose

    def log(self, message: str) -> None:
        if self.verbose:
            print(message, flush=True)

    def request_json(self, url: str, params: Dict[str, object]) -> Dict:
        last_error: Optional[Exception] = None
        for attempt in range(1, 4):
            try:
                response = self.session.get(url, params=params, timeout=self.timeout)
                response.raise_for_status()
                data = response.json()
                time.sleep(self.delay)
                return data
            except Exception as exc:
                last_error = exc
                wait_seconds = 2 * attempt
                self.log(f"[WARN] 请求失败，第 {attempt} 次重试，等待 {wait_seconds}s：{exc}")
                time.sleep(wait_seconds)
        raise RuntimeError(f"请求失败，已重试 3 次：{last_error}")

    def get_category_pages(self, category: str, limit: int = 200) -> List[str]:
        api_url = "https://en.wikipedia.org/w/api.php"
        pages: List[str] = []
        continuation: Dict[str, object] = {}
        while len(pages) < limit:
            params = {
                "action": "query",
                "format": "json",
                "list": "categorymembers",
                "cmtitle": category,
                "cmnamespace": 0,
                "cmlimit": min(50, limit - len(pages)),
                **continuation,
            }
            data = self.request_json(api_url, params)
            members = data.get("query", {}).get("categorymembers", [])
            pages.extend([item["title"] for item in members if "title" in item])
            if "continue" not in data:
                break
            continuation = data["continue"]
        return pages

    def get_page_info(self, title: str) -> Tuple[str, Optional[str], str]:
        api_url = "https://en.wikipedia.org/w/api.php"
        params = {
            "action": "query",
            "format": "json",
            "prop": "extracts|pageprops",
            "titles": title,
            "exintro": 1,
            "explaintext": 1,
            "redirects": 1,
        }
        data = self.request_json(api_url, params)
        pages = data.get("query", {}).get("pages", {})
        if not pages:
            return "", None, title
        page = next(iter(pages.values()))
        extract = (page.get("extract") or "").strip()
        qid = page.get("pageprops", {}).get("wikibase_item")
        normalized_title = page.get("title", title)
        return extract, qid, normalized_title

    def get_wikidata_info(self, qid: Optional[str]) -> Dict[str, str]:
        if not qid:
            return {}
        api_url = "https://www.wikidata.org/w/api.php"
        params = {
            "action": "wbgetentities",
            "format": "json",
            "ids": qid,
            "props": "labels|descriptions|sitelinks|aliases",
            "languages": "zh-hans|zh-cn|zh|zh-hant|en|ru",
            "sitefilter": "zhwiki|enwiki|ruwiki",
        }
        data = self.request_json(api_url, params)
        entity = data.get("entities", {}).get(qid, {})
        labels = entity.get("labels", {})
        descriptions = entity.get("descriptions", {})
        sitelinks = entity.get("sitelinks", {})

        def label(lang: str) -> str:
            return labels.get(lang, {}).get("value", "").strip()

        def desc(lang: str) -> str:
            return descriptions.get(lang, {}).get("value", "").strip()

        def sitelink(site: str) -> str:
            return sitelinks.get(site, {}).get("title", "").strip()

        return {
            "zh": label("zh-hans") or label("zh-cn") or label("zh") or label("zh-hant"),
            "en": label("en"),
            "ru": label("ru"),
            "zh_desc": desc("zh-hans") or desc("zh-cn") or desc("zh") or desc("zh-hant"),
            "en_desc": desc("en"),
            "ru_desc": desc("ru"),
            "zhwiki_title": sitelink("zhwiki"),
            "enwiki_title": sitelink("enwiki"),
            "ruwiki_title": sitelink("ruwiki"),
        }

    @staticmethod
    def valid_term_title(title: str) -> bool:
        title = title.strip()
        bad_prefixes = ("List of ", "Lists of ", "History of ", "Timeline of ", "Index of ", "Outline of ", "Glossary of ", "Comparison of ")
        if not title:
            return False
        if any(title.startswith(prefix) for prefix in bad_prefixes):
            return False
        if "(disambiguation)" in title:
            return False
        if len(title) > 100:
            return False
        if title.count(",") >= 2:
            return False
        return True

    @staticmethod
    def first_sentences(text: str, max_chars: int = 500) -> str:
        text = clean_text(text)
        if len(text) <= max_chars:
            return text
        candidates = re.split(r"(?<=[.!?。！？])\s+", text)
        result = ""
        for sentence in candidates:
            if len(result) + len(sentence) <= max_chars:
                result = (result + " " + sentence).strip()
            else:
                break
        return result or text[:max_chars].strip()

    @staticmethod
    def wiki_url(lang: str, title: str) -> str:
        return f"https://{lang}.wikipedia.org/wiki/{quote(title.replace(' ', '_'))}"

    def make_term_row(
        self,
        zh_term: str,
        en_term: str,
        ru_term: str,
        extract: str,
        wd: Dict[str, str],
        category: str,
        qid: Optional[str],
        normalized_title: str,
        original_title: str,
        status: str,
        contributor_name: str,
        contributor_research_area: str,
    ) -> TermRow:
        en_desc = wd.get("en_desc", "")
        zh_desc = wd.get("zh_desc", "")
        academic_explanation = self.first_sentences(extract or zh_desc or en_desc, max_chars=900)
        short_en_explanation = self.first_sentences(en_desc or extract, max_chars=500)
        if zh_desc and not zh_term.startswith("待译："):
            plain_explanation = f"{zh_term}通常指{zh_desc}。"
        else:
            plain_explanation = self.first_sentences(academic_explanation, max_chars=350) if academic_explanation else ""

        refs = []
        enwiki_title = wd.get("enwiki_title") or normalized_title or original_title
        refs.append(self.wiki_url("en", enwiki_title))
        if wd.get("zhwiki_title"):
            refs.append(self.wiki_url("zh", wd["zhwiki_title"]))
        if qid:
            refs.append(f"https://www.wikidata.org/wiki/{qid}")

        return TermRow(
            zh_term=zh_term,
            en_term=en_term,
            ru_term=ru_term,
            academic_explanation=academic_explanation,
            plain_explanation=plain_explanation,
            short_en_explanation=short_en_explanation,
            keywords=category_to_keywords(category),
            related_terms="",
            application_scenario=category_to_scenario(category),
            formula_or_model="",
            has_video_case="否",
            references="; ".join(refs[:3]),
            contributor_name=contributor_name,
            contributor_research_area=contributor_research_area,
            status=status,
        )

    def make_skipped_row(self, reason: str, en_term: str, ru_term: str, extract: str, wd: Dict[str, str], category: str, qid: Optional[str], normalized_title: str, original_title: str) -> SkippedRow:
        enwiki_title = wd.get("enwiki_title") or normalized_title or original_title
        return SkippedRow(
            action="建议人工翻译中文术语后导入",
            reason=reason,
            en_term=en_term,
            suggested_zh_term=f"待译：{en_term}",
            ru_term=ru_term,
            english_summary=self.first_sentences(wd.get("en_desc") or extract, max_chars=700),
            source_category=category,
            qid=qid or "",
            enwiki_url=self.wiki_url("en", enwiki_title),
            wikidata_url=f"https://www.wikidata.org/wiki/{qid}" if qid else "",
        )

    def build_rows(self, categories: Iterable[str], target: int, per_category_limit: int, status: str, contributor_name: str, contributor_research_area: str, missing_zh_mode: str) -> Tuple[List[TermRow], List[SkippedRow]]:
        seen_titles: Set[str] = set()
        seen_terms: Set[str] = set()
        rows: List[TermRow] = []
        skipped_rows: List[SkippedRow] = []

        for category in categories:
            if len(rows) >= target and missing_zh_mode in {"skip", "importable"}:
                break
            self.log(f"\n[INFO] 抓取分类：{category}")
            try:
                titles = self.get_category_pages(category, limit=per_category_limit)
            except Exception as exc:
                self.log(f"[ERROR] 分类抓取失败：{category}：{exc}")
                continue
            self.log(f"[INFO] 分类获得页面数：{len(titles)}")

            for title in titles:
                if len(rows) >= target and missing_zh_mode in {"skip", "importable"}:
                    break
                if title in seen_titles:
                    continue
                seen_titles.add(title)
                if not self.valid_term_title(title):
                    continue

                try:
                    extract, qid, normalized_title = self.get_page_info(title)
                    wd = self.get_wikidata_info(qid)
                except Exception as exc:
                    self.log(f"[WARN] 跳过 {title}，原因：{exc}")
                    continue

                en_term = wd.get("en") or normalized_title or title
                zh_term = wd.get("zh") or wd.get("zhwiki_title") or ""
                ru_term = wd.get("ru") or wd.get("ruwiki_title") or ""

                term_key = f"{zh_term.lower()}|{en_term.lower()}" if zh_term else f"missingzh|{en_term.lower()}"
                if term_key in seen_terms:
                    continue
                seen_terms.add(term_key)

                if not zh_term:
                    skipped = self.make_skipped_row(
                        reason="缺少 Wikidata 中文 label / 中文维基链接",
                        en_term=en_term,
                        ru_term=ru_term,
                        extract=extract,
                        wd=wd,
                        category=category,
                        qid=qid,
                        normalized_title=normalized_title,
                        original_title=title,
                    )
                    if missing_zh_mode in {"separate", "both"}:
                        skipped_rows.append(skipped)
                    if missing_zh_mode in {"importable", "both"} and len(rows) < target:
                        row = self.make_term_row(
                            zh_term=skipped.suggested_zh_term,
                            en_term=en_term,
                            ru_term=ru_term,
                            extract=extract,
                            wd=wd,
                            category=category,
                            qid=qid,
                            normalized_title=normalized_title,
                            original_title=title,
                            status=status,
                            contributor_name=contributor_name,
                            contributor_research_area=contributor_research_area,
                        )
                        rows.append(row)
                        self.log(f"[WAIT-ZH] {len(rows):03d}/{target} {row.zh_term} / {en_term}")
                    else:
                        self.log(f"[SKIP→SHEET] 缺少中文术语：{title}")
                    continue

                if len(rows) >= target:
                    continue
                row = self.make_term_row(zh_term, en_term, ru_term, extract, wd, category, qid, normalized_title, title, status, contributor_name, contributor_research_area)
                rows.append(row)
                self.log(f"[OK] {len(rows):03d}/{target} {zh_term} / {en_term}")

        return rows, skipped_rows


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"\s+", " ", text)
    text = text.replace("\u200b", "")
    return text.strip()


def category_to_keywords(category: str) -> str:
    category_lower = category.lower()
    if "photovoltaic" in category_lower:
        return "光伏, 太阳能, 可再生能源, 发电技术"
    if "solar" in category_lower:
        return "太阳能, 可再生能源, 能源转换, 清洁能源"
    if "storage" in category_lower:
        return "储能, 电池, 能源系统, 电力工程"
    if "climate" in category_lower:
        return "碳中和, 气候变化, 减排, 可持续发展"
    if "carbon" in category_lower:
        return "碳中和, 碳排放, 温室气体, 减排"
    if "greenhouse" in category_lower:
        return "温室气体, 气候变化, 碳排放, 环境科学"
    if "emission" in category_lower:
        return "减排, 碳排放, 环境政策, 碳中和"
    return "能源, 可再生能源, 碳中和, 工程技术"


def category_to_scenario(category: str) -> str:
    category_lower = category.lower()
    if "photovoltaic" in category_lower:
        return "光伏发电系统、太阳能电池、分布式能源和新能源工程。"
    if "solar" in category_lower:
        return "太阳能利用、清洁能源开发、建筑能源系统和可再生能源工程。"
    if "storage" in category_lower:
        return "电网调峰、可再生能源并网、储能系统设计和能源管理。"
    if "climate" in category_lower:
        return "气候变化减缓、碳中和路径、能源政策和低碳技术评估。"
    if "carbon" in category_lower:
        return "碳排放核算、碳捕集利用与封存、碳市场和低碳能源系统。"
    if "greenhouse" in category_lower:
        return "温室气体排放评估、环境监测、气候模型和减排政策。"
    if "emission" in category_lower:
        return "污染物与温室气体减排、能源系统优化和环境政策分析。"
    return "能源工程、可再生能源系统、碳中和技术和专业术语学习。"


def style_sheet(ws, header_color: str) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor=header_color)
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def save_excel(rows: List[TermRow], skipped_rows: List[SkippedRow], output_path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EnerTri导入_维基百科"
    ws.append(COLUMNS)
    style_sheet(ws, "D9EAF7")
    for row in rows:
        ws.append(row.to_list())

    main_widths = {1: 20, 2: 26, 3: 24, 4: 60, 5: 45, 6: 55, 7: 28, 8: 24, 9: 42, 10: 28, 11: 20, 12: 70, 13: 24, 14: 30, 15: 14}
    for idx, width in main_widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if skipped_rows:
        skip_ws = wb.create_sheet("待人工翻译_SKIP")
        skip_ws.append(SKIPPED_COLUMNS)
        style_sheet(skip_ws, "FCE4D6")
        for row in skipped_rows:
            skip_ws.append(row.to_list())
        widths = [22, 30, 34, 34, 24, 70, 34, 16, 70, 45]
        for idx, width in enumerate(widths, start=1):
            skip_ws.column_dimensions[get_column_letter(idx)].width = width
        for row_cells in skip_ws.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        skip_ws.freeze_panes = "A2"
        skip_ws.auto_filter.ref = skip_ws.dimensions

    wb.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="从 Wikipedia / Wikidata 抓取能源术语并导出为 EnerTri 可导入 Excel。")
    parser.add_argument("--target", type=int, default=50, help="目标导出术语数量，默认 50。")
    parser.add_argument("--per-category-limit", type=int, default=200, help="每个 Wikipedia 分类最多读取多少个页面标题，默认 200。")
    parser.add_argument("--output", default="EnerTri_维基百科能源术语_可导入.xlsx", help="输出 Excel 文件名。")
    parser.add_argument("--status", default="submitted", choices=["submitted", "approved"], help="导入后的审核状态，默认 submitted。")
    parser.add_argument("--delay", type=float, default=0.5, help="每次请求之间的间隔秒数，默认 0.5。建议不要低于 0.3。")
    parser.add_argument("--user-agent", default="EnerTriTermBuilder/1.0 (educational term database; contact: example@example.com)", help="请求 Wikimedia API 使用的 User-Agent。建议改成你的项目名和邮箱。")
    parser.add_argument("--category", action="append", dest="categories", help="指定 Wikipedia 分类。可以多次使用。")
    parser.add_argument("--contributor-name", default="Wikipedia 自动采集", help="Excel 中“填写本词条的学生姓名”列的默认值。")
    parser.add_argument("--research-area", default="能源与碳中和术语库建设", help="Excel 中“填写本词条的学生研究方向”列的默认值。")
    parser.add_argument("--missing-zh-mode", choices=["skip", "separate", "importable", "both"], default="separate", help="缺少中文术语时如何处理：skip=仍然跳过；separate=放到第二个工作表；importable=用“待译：英文术语”生成可导入行；both=两者都做。默认 separate。")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    categories = args.categories or DEFAULT_CATEGORIES

    print("=" * 72)
    print("EnerTri Wikipedia 术语抓取脚本 V2")
    print("=" * 72)
    print(f"目标数量：{args.target}")
    print(f"输出文件：{args.output}")
    print(f"审核状态：{args.status}")
    print(f"请求间隔：{args.delay}s")
    print(f"缺中文处理：{args.missing_zh_mode}")
    print("抓取分类：")
    for category in categories:
        print(f"  - {category}")
    print("=" * 72)

    builder = WikiEnerTriBuilder(user_agent=args.user_agent, delay=args.delay, verbose=True)
    rows, skipped_rows = builder.build_rows(
        categories=categories,
        target=args.target,
        per_category_limit=args.per_category_limit,
        status=args.status,
        contributor_name=args.contributor_name,
        contributor_research_area=args.research_area,
        missing_zh_mode=args.missing_zh_mode,
    )

    if not rows and not skipped_rows:
        print("[ERROR] 没有生成任何术语。请检查网络、分类名称或 User-Agent。", file=sys.stderr)
        return 1

    save_excel(rows, skipped_rows, args.output)
    print("\n" + "=" * 72)
    print(f"完成：主导入表 {len(rows)} 条")
    print(f"待人工翻译表 {len(skipped_rows)} 条")
    print(f"Excel 文件：{args.output}")
    print("下一步：")
    print("  1. 主导入表可以直接上传到 EnerTri。")
    print("  2. 待人工翻译_SKIP 表建议补中文术语后再复制到主导入表。")
    print("=" * 72)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
