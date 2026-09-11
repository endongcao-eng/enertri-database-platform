"""从《数据库词条填写模板.xlsx》导入术语到当前 DATABASE_URL。

用法：
  cd backend
  python scripts/import_terms_from_excel.py ../../数据库词条填写模板.xlsx --status submitted
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.database import SessionLocal  # noqa: E402
from app.models import Category, Term  # noqa: E402
from app.migration_bootstrap import upgrade_database  # noqa: E402
from app.schemas import TermCreate  # noqa: E402
from app.security import hash_password  # noqa: E402
from app.utils import (  # noqa: E402
    apply_related_term_names,
    apply_term_payload,
    bool_video,
    build_term_name_lookup,
    seed_defaults,
    split_list,
    split_refs,
)

HEADERS = {
    "中文术语": "zh",
    "英文术语": "en",
    "俄文术语（可选）": "ru",
    "学术版术语解释": "definition_zh_academic",
    "通俗版术语解释": "definition_zh_popular",
    "简要英文解释": "definition_en",
    "关键词（3–5个，用逗号分隔）": "keywords",
    "相关术语": "related_terms_text",
    "应用场景": "application_scenario",
    "公式 / 模型（如有）": "formula_model",
    "有无视频形式的示例或工程案例": "has_video_text",
    "参考文献（1–2篇）": "references",
    "填写本词条的学生姓名": "contributor_name",
    "填写本词条的学生研究方向": "contributor_research_direction",
    "审核状态": "review_status",
}


def find_import_sheet(workbook):
    candidates = [workbook.active, *[sheet for sheet in workbook.worksheets if sheet is not workbook.active]]
    for sheet in candidates:
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        indexes = {HEADERS[header]: i for i, header in enumerate(headers) if header in HEADERS}
        if "zh" in indexes and "en" in indexes:
            return sheet, indexes
    raise ValueError("未找到可导入工作表：第一行必须包含“中文术语”和“英文术语”")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--category-code", default="DEFAULT", help="导入词条默认分类代码")
    parser.add_argument("--category-name", default="未分类", help="导入词条默认分类中文名")
    parser.add_argument("--status", default=None, choices=["draft", "submitted", "approved"])
    parser.add_argument("--skip-duplicates", action="store_true", help="按中文+英文跳过重复术语")
    args = parser.parse_args()

    upgrade_database("head")
    workbook = load_workbook(args.xlsx_path, data_only=True)
    worksheet, indexes = find_import_sheet(workbook)

    imported = skipped = 0
    errors: list[str] = []
    pending_relations: list[tuple[Term, int, list[str]]] = []
    seen_pairs: set[tuple[str, str]] = set()

    with SessionLocal() as db:
        seed_defaults(db, hash_password)
        default_cat = db.scalar(select(Category).where(Category.code == args.category_code))
        if not default_cat:
            default_cat = Category(code=args.category_code, name_zh=args.category_name, sort_order=999)
            db.add(default_cat)
            db.flush()

        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            data = {key: row[idx] for key, idx in indexes.items() if idx < len(row)}
            zh = str(data.get("zh") or "").strip()
            en = str(data.get("en") or "").strip()
            if not zh and not en:
                skipped += 1
                continue
            if not zh or not en:
                errors.append(f"第 {row_number} 行缺少中文术语或英文术语")
                skipped += 1
                continue

            pair = (zh, en)
            if pair in seen_pairs:
                errors.append(f"第 {row_number} 行重复：{zh} / {en}，已跳过")
                skipped += 1
                continue
            seen_pairs.add(pair)

            existing = db.scalar(select(Term).where(Term.zh == zh, Term.en == en))
            if existing and args.skip_duplicates:
                skipped += 1
                continue

            status = (args.status or str(data.get("review_status") or "draft")).strip().lower()
            if status not in {"draft", "submitted", "approved"}:
                status = "draft"
            payload = TermCreate(
                category_id=default_cat.id,
                zh=zh,
                en=en,
                ru=str(data.get("ru") or "").strip() or None,
                definition_zh_academic=str(data.get("definition_zh_academic") or "待补充").strip(),
                definition_zh_popular=str(data.get("definition_zh_popular") or "").strip() or None,
                definition_en=str(data.get("definition_en") or "").strip() or None,
                application_scenario=str(data.get("application_scenario") or "").strip() or None,
                formula_model=str(data.get("formula_model") or "").strip() or None,
                has_video=bool_video(data.get("has_video_text")),
                review_status=status,
                contributor_name=str(data.get("contributor_name") or "").strip() or None,
                contributor_research_direction=str(data.get("contributor_research_direction") or "").strip() or None,
                keywords=split_list(data.get("keywords")),
                references=split_refs(data.get("references")),
                video_path=None,
            )
            term = existing or Term()
            apply_term_payload(db, term, payload)
            db.add(term)
            pending_relations.append((term, row_number, split_list(data.get("related_terms_text"))))
            imported += 1

        db.flush()
        lookup = build_term_name_lookup(db.scalars(select(Term)).all())
        for source_term, row_number, related_names in pending_relations:
            for warning in apply_related_term_names(source_term, related_names, lookup):
                errors.append(f"第 {row_number} 行：{warning}")
        db.commit()

    print({"imported": imported, "skipped": skipped, "errors": errors})


if __name__ == "__main__":
    main()
