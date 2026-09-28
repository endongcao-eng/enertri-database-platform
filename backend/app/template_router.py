from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from .audit import write_audit
from .database import MEDIA_DIR, get_db
from .file_service import FILE_STORE_DIR, InvalidStoredUpload, store_upload
from .models import TemplateFile, User, WorkspaceFile
from .security import current_user, require_permission

router = APIRouter(prefix="/api/templates", tags=["Template Library"])

CATEGORY_LABELS = {"pdf": "PDF", "ppt": "PPT", "excel": "EXCEL", "word": "WORD"}
CATEGORY_SUFFIXES = {
    "pdf": {".pdf"},
    "ppt": {".ppt", ".pptx"},
    "excel": {".xls", ".xlsx"},
    "word": {".doc", ".docx"},
}
TEMPLATE_SUFFIXES = {suffix for values in CATEGORY_SUFFIXES.values() for suffix in values}


def _meta(request: Request) -> dict[str, str | None]:
    return {"ip_address": request.client.host if request.client else None, "user_agent": request.headers.get("user-agent")}


def _template_to_dict(item: TemplateFile) -> dict:
    file = item.file
    return {
        "id": item.id,
        "file_id": file.id,
        "title": item.title,
        "category": item.category,
        "category_label": CATEGORY_LABELS[item.category],
        "original_name": file.original_name,
        "size_bytes": file.size_bytes,
        "mime_type": file.mime_type,
        "created_at": item.created_at,
        "uploader_name": item.uploader.display_name if item.uploader else "",
        "download_url": f"/api/templates/{item.id}/download",
    }


def _template_path(item: TemplateFile) -> Path:
    if not item or not item.file or item.file.deleted_at is not None:
        raise HTTPException(status_code=404, detail="模板不存在")
    path = Path(item.file.storage_path).resolve()
    media_root = MEDIA_DIR.resolve()
    file_store_root = FILE_STORE_DIR.resolve()
    allowed_roots = (media_root, file_store_root)
    inside_allowed_root = any(path != root and root in path.parents for root in allowed_roots)
    if not path.exists() or not inside_allowed_root:
        raise HTTPException(status_code=404, detail="模板文件内容不存在")
    return path


@router.get("")
def list_templates(
    category: str | None = Query(default=None),
    q: str | None = Query(default=None, max_length=200),
    limit: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission("template.view")),
):
    if category and category not in CATEGORY_LABELS:
        raise HTTPException(status_code=400, detail="模板分类必须是 PDF、PPT、EXCEL 或 WORD")
    stmt = select(TemplateFile).join(TemplateFile.file).where(WorkspaceFile.deleted_at.is_(None)).order_by(TemplateFile.id.desc()).limit(limit)
    if category:
        stmt = stmt.where(TemplateFile.category == category)
    if q:
        like = f"%{q.strip()}%"
        stmt = stmt.where(or_(TemplateFile.title.ilike(like), WorkspaceFile.original_name.ilike(like)))
    return [_template_to_dict(item) for item in db.scalars(stmt).all()]


@router.post("")
def upload_template(
    request: Request,
    file: UploadFile = File(...),
    category: str = Form(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("template.manage")),
):
    if category not in CATEGORY_LABELS:
        raise HTTPException(status_code=400, detail="模板分类必须是 PDF、PPT、EXCEL 或 WORD")
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in CATEGORY_SUFFIXES[category]:
        raise HTTPException(status_code=400, detail=f"{CATEGORY_LABELS[category]} 分类仅支持 {', '.join(sorted(CATEGORY_SUFFIXES[category]))} 文件")
    try:
        stored, duplicate = store_upload(
            db, user, file, confidentiality="internal", external_ai_allowed=False,
            allowed_suffixes=TEMPLATE_SUFFIXES, parse_status="not_applicable",
        )
    except InvalidStoredUpload as exc:
        write_audit(db, action="template.upload.rejected", resource_type="template_file", resource_id=exc.record.id, actor=user, outcome="failed", details=exc.validation, **_meta(request))
        raise
    existing = db.scalar(select(TemplateFile).where(TemplateFile.file_id == stored.id))
    if existing:
        if existing.category != category:
            raise HTTPException(status_code=409, detail="该文件已存在于其他模板分类")
        write_audit(db, action="template.duplicate_detected", resource_type="template_file", resource_id=existing.id, actor=user, details={"file_id": stored.id}, **_meta(request))
        return _template_to_dict(existing) | {"duplicate_detected": True}
    item = TemplateFile(file_id=stored.id, category=category, title=stored.original_name, uploader_id=user.id)
    db.add(item)
    db.commit()
    db.refresh(item)
    write_audit(db, action="template.uploaded", resource_type="template_file", resource_id=item.id, actor=user, details={"file_id": stored.id, "category": category, "duplicate": duplicate}, **_meta(request))
    return _template_to_dict(item) | {"duplicate_detected": duplicate}


@router.get("/{template_id}/download")
def download_template(template_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require_permission("template.download"))):
    item = db.scalar(select(TemplateFile).where(TemplateFile.id == template_id).join(TemplateFile.file))
    path = _template_path(item)
    write_audit(db, action="template.downloaded", resource_type="template_file", resource_id=item.id, actor=user, **_meta(request))
    return FileResponse(path, filename=item.file.original_name, media_type=item.file.mime_type or "application/octet-stream")


@router.get("/{template_id}/preview")
def preview_template(template_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require_permission("template.view"))):
    item = db.scalar(select(TemplateFile).where(TemplateFile.id == template_id).join(TemplateFile.file))
    path = _template_path(item)
    suffix = path.suffix.lower()
    result: dict = {"id": item.id, "title": item.title, "category": item.category, "original_name": item.file.original_name, "kind": "text", "sections": []}
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path), strict=False)
            result["kind"] = "pdf"
            result["page_count"] = len(reader.pages)
            result["pdf_url"] = f"/api/templates/{item.id}/download"
            result["sections"] = [{"title": f"第 {index + 1} 页", "text": (page.extract_text() or "").strip()[:12000]} for index, page in enumerate(reader.pages[:20])]
        elif suffix == ".docx":
            from docx import Document
            document = Document(str(path))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            tables = [[[cell.text.strip() for cell in row.cells] for row in table.rows[:50]] for table in document.tables[:10]]
            result["kind"] = "docx"
            result["sections"] = [{"title": "正文", "text": "\n".join(paragraphs[:300])}]
            result["tables"] = tables
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(str(path))
            result["kind"] = "pptx"
            result["slide_count"] = len(presentation.slides)
            result["sections"] = [{"title": f"第 {index + 1} 页", "text": "\n".join(shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip())[:12000]} for index, slide in enumerate(presentation.slides[:30])]
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            result["kind"] = "xlsx"
            result["sheets"] = []
            for sheet in workbook.worksheets[:10]:
                rows = []
                for row in sheet.iter_rows(max_row=80, max_col=30, values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(values): rows.append(values)
                result["sheets"].append({"title": sheet.title, "rows": rows})
        else:
            result["kind"] = "unsupported"
            result["message"] = "旧版 Office 文件暂不支持网页内结构化预览，请下载后使用本地 Office 软件打开。"
    except Exception as exc:
        result["kind"] = "unsupported"
        result["message"] = f"无法生成网页预览，请下载后查看：{exc}"
    write_audit(db, action="template.previewed", resource_type="template_file", resource_id=item.id, actor=user, details={"kind": result.get("kind")}, **_meta(request))
    return result


@router.delete("/{template_id}")
def delete_template(template_id: int, request: Request, db: Session = Depends(get_db), user: User = Depends(require_permission("template.delete"))):
    item = db.get(TemplateFile, template_id)
    if not item:
        raise HTTPException(status_code=404, detail="模板不存在")
    file = item.file
    if file:
        file.deleted_at = datetime.now(timezone.utc)
    db.delete(item)
    db.commit()
    write_audit(db, action="template.deleted", resource_type="template_file", resource_id=template_id, actor=user, **_meta(request))
    return {"deleted": template_id}
