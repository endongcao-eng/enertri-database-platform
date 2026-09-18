from __future__ import annotations

import os
import random
import secrets
import shutil
import time
from datetime import datetime, timezone
from pathlib import Path

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy import or_, select
from sqlalchemy.orm import Session, selectinload

from .database import MEDIA_DIR, SessionLocal, get_db
from .models import Article, Category, Term, User, UserTermProgress
from .schemas import ArticleCreate, ArticleOut, ArticleUpdate, CategoryCreate, CategoryOut, ImportResult, LoginRequest, ProgressOut, ProgressUpdate, QuizOut, QuizSubmitRequest, QuizSubmitResult, ReviewRequest, TermCreate, TermOut, TermUpdate, TokenOut, UserOut
from .security import admin_user, create_access_token, current_user, hash_password, optional_current_user, verify_password, require_permission
from .utils import apply_related_term_names, article_to_out, apply_term_payload, bool_video, build_term_name_lookup, join_keywords, progress_to_out, seed_defaults, split_list, split_refs, term_to_out
from .workspace_router import router as workspace_router
from .foundation_router import router as foundation_router
from .knowledge_router import router as knowledge_router
from .capability_router import router as capability_router
from .research_profile_router import router as research_profile_router
from .project_matching_router import router as project_matching_router
from .development_planning_router import router as development_planning_router
from .development_execution_router import router as development_execution_router
from .capability_standard_service import seed_capability_standards
from .research_profile_service import seed_demo_research_profiles
from .project_matching_service import seed_demo_project_matching
from .development_planning_service import seed_demo_development_planning
from .development_execution_service import seed_demo_closed_loop
from .migration_bootstrap import upgrade_database
from .file_service import cleanup_expired_files
from .file_validation import FileValidationError, validate_file_bytes
from .audit import write_audit

app = FastAPI(title="EnerTri Research & Welding Platform API", version="4.7.0", docs_url="/api/docs", openapi_url="/api/openapi.json")

# 生产环境请在 .env 中设置 CORS_ORIGINS=https://your-domain.com
# 多个来源用英文逗号分隔；开发阶段可保留本地地址。
cors_origins = [origin.strip() for origin in os.getenv("CORS_ORIGINS", "http://127.0.0.1:5500,http://localhost:5500").split(",") if origin.strip()]

app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/media/videos", StaticFiles(directory=str(MEDIA_DIR / "videos")), name="public_videos")
from .account_router import router as account_router
app.include_router(account_router)
app.include_router(workspace_router)
app.include_router(foundation_router)
app.include_router(knowledge_router)
app.include_router(capability_router)
app.include_router(research_profile_router)
app.include_router(project_matching_router)
app.include_router(development_planning_router)
app.include_router(development_execution_router)


@app.on_event("startup")
def startup_checks() -> None:
    # V4.2 production deployments must use the one-shot init service.
    # This opt-in path exists only for isolated local/test runs.
    if os.getenv("APP_BOOTSTRAP_ON_STARTUP", "false").lower() in {"1", "true", "yes", "on"}:
        from .task_queue import recover_expired_tasks
        migration_result = upgrade_database("head")
        with SessionLocal() as db:
            seed_defaults(db, hash_password)
            standard_seed = seed_capability_standards(db)
            profile_seed = seed_demo_research_profiles(db)
            matching_seed = seed_demo_project_matching(db, hash_password)
            planning_seed = seed_demo_development_planning(db)
            closed_loop_seed = seed_demo_closed_loop(db)
            interrupted = recover_expired_tasks()
            cleaned = cleanup_expired_files(db)
            write_audit(db, action="database.migration.completed", resource_type="database", resource_id="v4_7_trial_controls", details={**migration_result, "capability_standard_seed": standard_seed, "research_profile_seed": profile_seed, "project_matching_seed": matching_seed, "development_planning_seed": planning_seed, "closed_loop_seed": closed_loop_seed, "expired_task_leases_recovered": interrupted, "expired_files_cleaned": cleaned, "bootstrap_mode": "explicit_startup_opt_in"})


@app.get("/api/health")
def health():
    return {"ok": True, "version": "4.7.0", "task_execution_mode": os.getenv("TASK_EXECUTION_MODE", "database_worker")}


@app.get("/api/ready")
def readiness():
    from .migration_bootstrap import current_revision, HEAD_REVISION
    try:
        if current_revision() != HEAD_REVISION:
            raise RuntimeError("schema mismatch")
    except Exception:
        raise HTTPException(status_code=503, detail="Database is not ready")
    return {"ok": True}


@app.post("/api/auth/login", response_model=TokenOut)
def login(payload: LoginRequest, request: Request, db: Session = Depends(get_db)):
    from .controls import rate_limit
    rate_limit(["login:user:"+payload.username, "login:ip:"+(request.client.host if request.client else "unknown")], [int(os.getenv("LOGIN_ACCOUNT_LIMIT","5")),int(os.getenv("LOGIN_IP_LIMIT","120"))], int(os.getenv("LOGIN_WINDOW_SECONDS","900")))
    user = db.scalar(select(User).where(User.username == payload.username))
    if not user or not verify_password(payload.password, user.password_hash):
        write_audit(db, action="auth.login.failed", resource_type="user", actor_user_id=user.id if user else None, outcome="failed")
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="账户已停用")
    write_audit(db, action="auth.login.succeeded", resource_type="user", actor=user)
    token = create_access_token(subject=user.username, role=user.role, token_version=user.token_version)
    return TokenOut(access_token=token, user=UserOut.model_validate(user))


@app.get("/api/auth/me", response_model=UserOut)
def me(user: User = Depends(current_user)):
    return user


@app.get("/api/categories", response_model=list[CategoryOut])
def list_categories(db: Session = Depends(get_db)):
    return db.scalars(select(Category).order_by(Category.sort_order, Category.id)).all()


@app.post("/api/categories", response_model=CategoryOut)
def create_category(payload: CategoryCreate, db: Session = Depends(get_db), _: User = Depends(admin_user)):
    item = Category(**payload.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


@app.get("/api/articles", response_model=list[ArticleOut])
def list_articles(
    q: str | None = Query(default=None),
    status: str | None = Query(default="approved", description="approved/draft/all；非管理员只能访问 approved"),
    limit: int = Query(default=50, le=200),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    user: User | None = Depends(optional_current_user),
):
    requested_status = status or "approved"
    if requested_status not in {"approved", "draft", "all"}:
        raise HTTPException(status_code=400, detail="无效的文章状态")
    if (not user or user.role not in {"admin", "system_admin"}) and requested_status != "approved":
        raise HTTPException(status_code=403, detail="未发布文章仅管理员可查看")

    stmt = select(Article).order_by(Article.id.desc()).offset(offset).limit(limit)
    if requested_status != "all":
        stmt = stmt.where(Article.status == requested_status)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(Article.title.like(like), Article.summary.like(like), Article.content.like(like), Article.related_keywords.like(like)))
    return [article_to_out(a) for a in db.scalars(stmt).all()]


@app.get("/api/articles/{article_id}", response_model=ArticleOut)
def get_article(
    article_id: int,
    db: Session = Depends(get_db),
    user: User | None = Depends(optional_current_user),
):
    article = db.get(Article, article_id)
    if not article or (article.status != "approved" and (not user or user.role not in {"admin", "system_admin"})):
        raise HTTPException(status_code=404, detail="article not found")
    return article_to_out(article)


@app.post("/api/articles", response_model=ArticleOut)
def create_article(payload: ArticleCreate, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    article = Article(
        title=payload.title,
        summary=payload.summary,
        content=payload.content,
        difficulty=payload.difficulty,
        status=payload.status,
        cover_label=payload.cover_label,
        related_keywords=join_keywords(payload.related_keywords),
        created_by=user.id,
        updated_by=user.id,
    )
    db.add(article)
    db.commit()
    db.refresh(article)
    return article_to_out(article)


@app.put("/api/articles/{article_id}", response_model=ArticleOut)
def update_article(article_id: int, payload: ArticleUpdate, db: Session = Depends(get_db), user: User = Depends(admin_user)):
    article = db.get(Article, article_id)
    if not article:
        raise HTTPException(status_code=404, detail="article not found")
    article.title = payload.title
    article.summary = payload.summary
    article.content = payload.content
    article.difficulty = payload.difficulty
    article.status = payload.status
    article.cover_label = payload.cover_label
    article.related_keywords = join_keywords(payload.related_keywords)
    article.updated_by = user.id
    db.commit()
    db.refresh(article)
    return article_to_out(article)


@app.delete("/api/articles/{article_id}")
def delete_article(article_id: int, db: Session = Depends(get_db), _: User = Depends(admin_user)):
    article = db.get(Article, article_id)
    if not article:
        raise HTTPException(status_code=404, detail="article not found")
    db.delete(article)
    db.commit()
    return {"deleted": article_id}



def _get_or_create_progress(db: Session, user_id: int, term_id: int) -> UserTermProgress:
    progress = db.scalar(select(UserTermProgress).where(UserTermProgress.user_id == user_id, UserTermProgress.term_id == term_id))
    if progress:
        return progress
    progress = UserTermProgress(user_id=user_id, term_id=term_id)
    db.add(progress)
    db.flush()
    return progress


from .quiz_store import save_quiz, consume_quiz


def _dedupe_options(values: list[str], fallback: list[str]) -> list[str]:
    seen: set[str] = set()
    output: list[str] = []
    for value in values + fallback:
        value = (value or "").strip()
        if value and value not in seen:
            seen.add(value)
            output.append(value)
        if len(output) == 4:
            break
    return output


def _make_quiz_question(
    question_id: str,
    question: str,
    correct_text: str,
    distractors: list[str],
    fallback: list[str],
    explanation: str,
) -> tuple[dict, dict]:
    values = _dedupe_options([correct_text, *distractors], fallback)
    random.SystemRandom().shuffle(values)
    labels = ["A", "B", "C", "D"]
    options = [{"id": labels[i], "text": value} for i, value in enumerate(values)]
    correct_answer = next(option["id"] for option in options if option["text"] == correct_text)
    return (
        {"id": question_id, "question": question, "options": options},
        {"answer": correct_answer, "explanation": explanation},
    )


def build_term_quiz(term: Term, all_terms: list[Term]) -> QuizOut:
    other_terms = [t for t in all_terms if t.id != term.id]
    correct_definition = term.definition_zh_popular or term.definition_zh_academic
    term_keywords = [k.keyword for k in term.keywords]
    all_keywords = [k.keyword for t in all_terms for k in t.keywords if k.keyword]
    correct_keyword = term_keywords[0] if term_keywords else (term.category.name_zh if term.category else "能源")

    public_questions: list[dict] = []
    answer_key: dict[str, dict] = {}
    question_specs = [
        _make_quiz_question(
            "q1",
            f"“{term.zh}”对应的英文术语最可能是？",
            term.en,
            [t.en for t in other_terms[:8]],
            ["photovoltaic module", "inverter", "energy storage", "carbon neutrality"],
            f"{term.zh} 的英文术语为 {term.en}。",
        ),
        _make_quiz_question(
            "q2",
            f"下面哪一项最接近“{term.zh}”的通俗解释？",
            correct_definition,
            [(t.definition_zh_popular or t.definition_zh_academic) for t in other_terms[:8]],
            ["描述能源系统中设备、材料或过程的专业概念。", "与能源转换效率或运行场景有关的概念。", "用于解释低碳技术应用价值的基础知识。"],
            correct_definition,
        ),
        _make_quiz_question(
            "q3",
            f"下列哪个关键词最适合关联“{term.zh}”？",
            correct_keyword,
            all_keywords,
            ["能源", "效率", "储能", "电网"],
            f"该词条关键词包括：{', '.join(term_keywords) if term_keywords else correct_keyword}。",
        ),
    ]
    for public_question, private_answer in question_specs:
        public_questions.append(public_question)
        answer_key[public_question["id"]] = private_answer

    quiz_token = secrets.token_urlsafe(32)
    save_quiz(quiz_token, term.id, answer_key)
    return QuizOut(term_id=term.id, title=f"{term.zh} 随堂小测", quiz_token=quiz_token, questions=public_questions)


def base_term_query():
    return select(Term).options(
        selectinload(Term.category),
        selectinload(Term.contributor),
        selectinload(Term.keywords),
        selectinload(Term.references),
        selectinload(Term.media_assets),
        selectinload(Term.outgoing_relations),
    )



@app.get("/api/me/progress", response_model=list[ProgressOut])
def list_my_progress(db: Session = Depends(get_db), user: User = Depends(current_user)):
    rows = db.scalars(select(UserTermProgress).where(UserTermProgress.user_id == user.id).order_by(UserTermProgress.last_viewed_at.desc())).all()
    return [progress_to_out(row) for row in rows]


@app.post("/api/me/progress/{term_id}", response_model=ProgressOut)
def update_my_progress(term_id: int, payload: ProgressUpdate, db: Session = Depends(get_db), user: User = Depends(current_user)):
    term = db.get(Term, term_id)
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    progress = _get_or_create_progress(db, user.id, term_id)
    if payload.is_favorite is not None:
        progress.is_favorite = payload.is_favorite
    if payload.learned is not None:
        progress.learned_at = datetime.now(timezone.utc) if payload.learned else None
    if payload.note is not None:
        progress.note = payload.note
    db.commit()
    db.refresh(progress)
    return progress_to_out(progress)

@app.get("/api/terms", response_model=list[TermOut])
def list_terms(
    q: str | None = Query(default=None),
    category_id: int | None = Query(default=None),
    status: str | None = Query(default="approved", description="draft/submitted/approved/all；非管理员只能访问 approved"),
    limit: int = Query(default=200, le=1000),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    user: User | None = Depends(optional_current_user),
):
    requested_status = status or "approved"
    if requested_status not in {"draft", "submitted", "approved", "all"}:
        raise HTTPException(status_code=400, detail="无效的词条审核状态")
    if (not user or user.role not in {"admin", "system_admin"}) and requested_status != "approved":
        raise HTTPException(status_code=403, detail="草稿和待审核词条仅管理员可批量查看")

    stmt = base_term_query().order_by(Term.featured.desc(), Term.id.desc()).offset(offset).limit(limit)
    if category_id:
        stmt = stmt.where(Term.category_id == category_id)
    if requested_status != "all":
        stmt = stmt.where(Term.review_status == requested_status)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(Term.zh.like(like), Term.en.like(like), Term.ru.like(like), Term.definition_zh_academic.like(like)))
    return [term_to_out(t) for t in db.scalars(stmt).unique().all()]



@app.get("/api/terms/{term_id}/quiz", response_model=QuizOut)
def get_term_quiz(term_id: int, db: Session = Depends(get_db)):
    term = db.scalar(base_term_query().where(Term.id == term_id, Term.review_status == "approved"))
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    all_terms = db.scalars(base_term_query().where(Term.review_status == "approved").limit(100)).unique().all()
    return build_term_quiz(term, all_terms)


@app.post("/api/terms/{term_id}/quiz/submit", response_model=QuizSubmitResult)
def submit_term_quiz(
    term_id: int,
    payload: QuizSubmitRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(optional_current_user),
):
    term = db.scalar(base_term_query().where(Term.id == term_id, Term.review_status == "approved"))
    if not term:
        raise HTTPException(status_code=404, detail="term not found")

    answer_key = consume_quiz(payload.quiz_token, term_id)
    if answer_key is None:
        raise HTTPException(status_code=400, detail="测验已过期或与当前词条不匹配，请重新生成题目")

    results: list[dict] = []
    correct_count = 0
    for question_id, expected in answer_key.items():
        selected = payload.answers.get(question_id)
        is_correct = selected == expected["answer"]
        correct_count += int(is_correct)
        results.append({
            "id": question_id,
            "selected_answer": selected,
            "correct_answer": expected["answer"],
            "correct": is_correct,
            "explanation": expected["explanation"],
        })

    total = len(answer_key)
    score = round(correct_count / total * 100) if total else 0
    best_score = None
    if user:
        progress = _get_or_create_progress(db, user.id, term_id)
        progress.quiz_best_score = max(progress.quiz_best_score or 0, score)
        db.commit()
        best_score = progress.quiz_best_score
    return QuizSubmitResult(
        score=score,
        correct_count=correct_count,
        total=total,
        best_score=best_score,
        results=results,
    )


@app.get("/api/terms/{term_id}", response_model=TermOut)
def get_term(
    term_id: int,
    db: Session = Depends(get_db),
    user: User | None = Depends(optional_current_user),
):
    term = db.scalar(base_term_query().where(Term.id == term_id))
    can_view_private = bool(user and (user.role in {"admin", "system_admin"} or term and term.created_by == user.id))
    if not term or (term.review_status != "approved" and not can_view_private):
        raise HTTPException(status_code=404, detail="term not found")
    return term_to_out(term)


@app.post("/api/terms", response_model=TermOut)
def create_term(payload: TermCreate, db: Session = Depends(get_db), user: User = Depends(current_user)):
    if user.role not in {"admin", "system_admin"}:
        payload.review_status = "draft"
        payload.review_comment = None
        payload.featured = False
    term = apply_term_payload(db, Term(), payload, user_id=user.id)
    db.add(term)
    db.commit()
    term = db.scalar(base_term_query().where(Term.id == term.id))
    return term_to_out(term)


@app.put("/api/terms/{term_id}", response_model=TermOut)
def update_term(term_id: int, payload: TermUpdate, db: Session = Depends(get_db), user: User = Depends(current_user)):
    term = db.scalar(base_term_query().where(Term.id == term_id))
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    if user.role not in {"admin", "system_admin"} and term.created_by != user.id:
        raise HTTPException(status_code=403, detail="只能编辑自己创建的词条")
    if user.role not in {"admin", "system_admin"} and term.review_status != "draft":
        raise HTTPException(status_code=403, detail="已提交或已通过审核的词条不能由学生直接修改")
    if user.role not in {"admin", "system_admin"}:
        # 审核状态、审核意见和推荐标记只能由管理员改变。
        payload.review_status = term.review_status
        payload.review_comment = term.review_comment
        payload.featured = term.featured
    apply_term_payload(db, term, payload, user_id=user.id)
    db.commit()
    term = db.scalar(base_term_query().where(Term.id == term_id))
    return term_to_out(term)


@app.post("/api/terms/{term_id}/submit", response_model=TermOut)
def submit_term(term_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    term = db.scalar(base_term_query().where(Term.id == term_id))
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    if user.role not in {"admin", "system_admin"} and term.created_by != user.id:
        raise HTTPException(status_code=403, detail="只能提交自己创建的词条")
    if term.review_status != "draft":
        raise HTTPException(status_code=400, detail="只有草稿状态的词条可以提交审核")
    term.review_status = "submitted"
    db.commit()
    term = db.scalar(base_term_query().where(Term.id == term_id))
    return term_to_out(term)


@app.post("/api/terms/{term_id}/review", response_model=TermOut)
def review_term(term_id: int, payload: ReviewRequest, db: Session = Depends(get_db), _: User = Depends(require_permission("process.approve"))):
    term = db.scalar(base_term_query().where(Term.id == term_id))
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    term.review_status = payload.review_status
    term.review_comment = payload.review_comment
    db.commit()
    term = db.scalar(base_term_query().where(Term.id == term_id))
    return term_to_out(term)


@app.delete("/api/terms/{term_id}")
def delete_term(term_id: int, db: Session = Depends(get_db), _: User = Depends(admin_user)):
    term = db.get(Term, term_id)
    if not term:
        raise HTTPException(status_code=404, detail="term not found")
    db.delete(term)
    db.commit()
    return {"deleted": term_id}


@app.post("/api/media/videos")
def upload_video(file: UploadFile = File(...), _: User = Depends(current_user)):
    safe_name = Path(file.filename or "video.mp4").name
    raw = file.file.read(512 * 1024 * 1024 + 1)
    if len(raw) > 512 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="视频文件超过 512MB 限制")
    try:
        validation = validate_file_bytes(safe_name, raw)
    except FileValidationError as exc:
        raise HTTPException(status_code=400, detail=f"视频文件校验失败：{exc}") from exc
    if validation.get("category") != "video":
        raise HTTPException(status_code=400, detail="上传内容不是可识别的视频文件")
    target = MEDIA_DIR / "videos" / safe_name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(raw)
    return {"path": f"/media/videos/{safe_name}", "filename": safe_name, "validation": validation}


@app.post("/api/import/excel", response_model=ImportResult)
def import_excel(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(admin_user)):
    from openpyxl import load_workbook
    from tempfile import NamedTemporaryFile

    headers_map = {
        "中文术语": "zh",
        "英文术语": "en",
        "俄文术语（可选）": "ru",
        "学术版术语解释": "definition_zh_academic",
        "通俗版术语解释": "definition_zh_popular",
        "简要英文解释": "definition_en",
        "关键词（3–5个，用逗号分隔）": "keywords",
        "相关术语": "related_terms_text",
        "应用场景": "application_scenario",
        "公式 / 模型（如有）": "formula_model",
        "有无视频形式的示例或工程案例": "has_video_text",
        "参考文献（1–2篇）": "references",
        "填写本词条的学生姓名": "contributor_name",
        "填写本词条的学生研究方向": "contributor_research_direction",
        "审核状态": "review_status",
    }

    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="目前仅支持 .xlsx Excel 工作簿文件")

    raw = file.file.read(64 * 1024 * 1024 + 1)
    if len(raw) > 64 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Excel 文件超过 64MB 限制")
    try:
        validation = validate_file_bytes(file.filename or "import.xlsx", raw)
    except FileValidationError as exc:
        raise HTTPException(status_code=400, detail=f"Excel 文件结构校验失败：{exc}") from exc
    if validation.get("category") != "office":
        raise HTTPException(status_code=400, detail="上传内容不是有效的 XLSX 工作簿")

    tmp_path = None
    try:
        with NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp.write(raw)
            tmp_path = tmp.name

        wb = load_workbook(tmp_path, data_only=True)

        # Excel 打开时的“当前活动工作表”不一定是导入表。
        # 因此优先寻找第一行同时包含“中文术语”和“英文术语”的工作表。
        ws = None
        indexes: dict[str, int] = {}
        for candidate in [wb.active, *[s for s in wb.worksheets if s is not wb.active]]:
            raw_headers = [cell.value for cell in candidate[1]]
            headers = [str(h).strip() if h is not None else "" for h in raw_headers]
            candidate_indexes = {headers_map[h]: i for i, h in enumerate(headers) if h in headers_map}
            if "zh" in candidate_indexes and "en" in candidate_indexes:
                ws = candidate
                indexes = candidate_indexes
                break

        if ws is None:
            raise HTTPException(
                status_code=400,
                detail="未找到可导入工作表：第一行必须包含“中文术语”和“英文术语”两列表头",
            )

        default_cat = db.scalar(select(Category).where(Category.code == "DEFAULT"))
        if not default_cat:
            default_cat = Category(code="DEFAULT", name_zh="未分类", sort_order=999)
            db.add(default_cat)
            db.flush()

        imported = skipped = 0
        errors: list[str] = []
        seen_pairs: set[tuple[str, str]] = set()
        pending_relations: list[tuple[Term, int, list[str]]] = []

        for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            data = {key: row[idx] for key, idx in indexes.items() if idx < len(row)}

            zh = str(data.get("zh") or "").strip()
            en = str(data.get("en") or "").strip()

            if not zh and not en:
                skipped += 1
                continue
            if not zh or not en:
                errors.append(f"第 {row_number} 行缺少中文术语或英文术语")
                skipped += 1
                continue

            pair = (zh, en)
            if pair in seen_pairs:
                skipped += 1
                errors.append(f"第 {row_number} 行重复：{zh} / {en}，已跳过")
                continue
            seen_pairs.add(pair)

            exists = db.scalar(select(Term.id).where(Term.zh == zh, Term.en == en))
            if exists:
                skipped += 1
                errors.append(f"第 {row_number} 行已存在：{zh} / {en}，已跳过")
                continue

            review_status = str(data.get("review_status") or "draft").strip().lower()
            if review_status not in {"draft", "submitted", "approved"}:
                review_status = "draft"

            payload = TermCreate(
                category_id=default_cat.id,
                zh=zh,
                en=en,
                ru=str(data.get("ru") or "").strip() or None,
                definition_zh_academic=str(data.get("definition_zh_academic") or "待补充").strip(),
                definition_zh_popular=str(data.get("definition_zh_popular") or "").strip() or None,
                definition_en=str(data.get("definition_en") or "").strip() or None,
                application_scenario=str(data.get("application_scenario") or "").strip() or None,
                formula_model=str(data.get("formula_model") or "").strip() or None,
                has_video=bool_video(data.get("has_video_text")),
                review_status=review_status,
                contributor_name=str(data.get("contributor_name") or "").strip() or None,
                contributor_research_direction=str(data.get("contributor_research_direction") or "").strip() or None,
                keywords=split_list(data.get("keywords")),
                references=split_refs(data.get("references")),
                video_path=None,
            )
            term = apply_term_payload(db, Term(), payload, user_id=user.id)
            db.add(term)
            pending_relations.append((term, row_number, split_list(data.get("related_terms_text"))))
            imported += 1

        # 两阶段导入：先让所有新词条获得 ID，再按中文/英文/俄文名称或 ID 建立关系。
        db.flush()
        lookup = build_term_name_lookup(db.scalars(select(Term)).all())
        for source_term, row_number, related_names in pending_relations:
            for warning in apply_related_term_names(source_term, related_names, lookup):
                errors.append(f"第 {row_number} 行：{warning}")

        db.commit()
        if len(errors) > 200:
            errors = errors[:200] + ["错误/跳过记录过多，仅显示前 200 条"]
        return ImportResult(imported=imported, skipped=skipped, errors=errors)

    except HTTPException:
        raise
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Excel 导入失败：{exc}") from exc
    finally:
        if tmp_path:
            try:
                Path(tmp_path).unlink(missing_ok=True)
            except Exception:
                pass


# Railway's single web service does not have Caddy in front of it. Serve the
# same static frontend from FastAPI after all API routes have been registered;
# the /api and /media mounts above therefore keep precedence.
FRONTEND_DIR = Path(__file__).resolve().parents[2] / "frontend_api_demo"
if FRONTEND_DIR.exists():
    @app.get("/login", include_in_schema=False)
    def login_page():
        return FileResponse(FRONTEND_DIR / "index.html")

    @app.get("/register", include_in_schema=False)
    def register_page():
        return FileResponse(FRONTEND_DIR / "index.html")

    app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")
