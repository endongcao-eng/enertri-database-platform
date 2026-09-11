from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches as PptInches, Pt as PptPt


CJK_FONT = "Noto Sans CJK SC"
LATIN_FONT = "Arial"


def safe_filename(value: str, suffix: str) -> str:
    base = re.sub(r"[^\w\-\u4e00-\u9fff]+", "_", value.strip(), flags=re.UNICODE).strip("_") or "enertri"
    return f"{base[:60]}{suffix}"


def _set_docx_style_font(style: Any, size: float | None = None) -> None:
    style.font.name = LATIN_FONT
    if size is not None:
        style.font.size = Pt(size)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)


def _set_docx_run_font(run: Any, size: float | None = None) -> None:
    run.font.name = LATIN_FONT
    if size is not None:
        run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)


def _set_ppt_paragraph_font(
    paragraph: Any,
    size: float | None = None,
    bold: bool | None = None,
    color: RGBColor | None = None,
) -> None:
    paragraph.font.name = CJK_FONT
    if size is not None:
        paragraph.font.size = PptPt(size)
    if bold is not None:
        paragraph.font.bold = bold
    if color is not None:
        paragraph.font.color.rgb = color
    for run in paragraph.runs:
        run.font.name = CJK_FONT
        if size is not None:
            run.font.size = PptPt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color


def _add_markdownish(document: Document, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
        else:
            document.add_paragraph(line)


def create_docx(path: Path, title: str, content: str, metadata: dict[str, Any] | None = None) -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = document.styles
    for style_name, style_size in {
        "Normal": 10.5,
        "Title": 20,
        "Heading 1": 16,
        "Heading 2": 14,
        "Heading 3": 12,
        "List Bullet": 10.5,
        "List Number": 10.5,
    }.items():
        if style_name in styles:
            _set_docx_style_font(styles[style_name], style_size)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    _set_docx_run_font(run, 20)
    if metadata:
        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(" · ".join(f"{k}: {v}" for k, v in metadata.items() if v))
        meta_run.italic = True
        _set_docx_run_font(meta_run, 9.5)
    document.add_paragraph()
    _add_markdownish(document, content)
    document.add_paragraph()
    note = document.add_paragraph("说明：本文件由 EnerTri 智能工作台生成，正式提交或生产使用前应由领域专家核验。")
    note.runs[0].italic = True
    _set_docx_run_font(note.runs[0], 9)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def _split_slides_from_text(title: str, content: str) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    current_title = "研究概览"
    bullets: list[str] = []
    for raw in content.splitlines():
        line = raw.strip()
        if line.startswith(("# ", "## ", "### ")):
            if bullets:
                sections.append({"title": current_title, "bullets": bullets[:7]})
            current_title = line.lstrip("# ").strip()
            bullets = []
        elif line.startswith(("- ", "* ")):
            bullets.append(line[2:].strip())
        elif line and len(line) <= 120:
            bullets.append(line)
    if bullets:
        sections.append({"title": current_title, "bullets": bullets[:7]})
    if not sections:
        sections = [{"title": "核心内容", "bullets": [p.strip() for p in re.split(r"[。；\n]", content) if p.strip()][:7]}]
    return [{"title": title, "subtitle": "EnerTri 学术汇报"}, *sections[:14]]


def create_pptx(path: Path, title: str, content: str, slides: list[dict[str, Any]] | None = None) -> Path:
    """Create a restrained 16:9 academic deck with editable text and shapes."""
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    slide_data = slides or _split_slides_from_text(title, content)

    navy = RGBColor(22, 42, 68)
    blue = RGBColor(38, 112, 176)
    cyan = RGBColor(73, 168, 202)
    ink = RGBColor(28, 35, 45)
    muted = RGBColor(102, 113, 128)
    paper = RGBColor(247, 249, 252)
    white = RGBColor(255, 255, 255)
    line = RGBColor(214, 222, 232)

    for index, item in enumerate(slide_data):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = paper

        if index == 0:
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PptInches(0.24), prs.slide_height)
            accent.fill.solid()
            accent.fill.fore_color.rgb = blue
            accent.line.fill.background()

            kicker = slide.shapes.add_textbox(PptInches(0.95), PptInches(0.72), PptInches(10.8), PptInches(0.35))
            kp = kicker.text_frame.paragraphs[0]
            kp.text = "ENERTRI · RESEARCH & WELDING INTELLIGENCE"
            _set_ppt_paragraph_font(kp, 11, True, blue)

            title_box = slide.shapes.add_textbox(PptInches(0.95), PptInches(2.05), PptInches(11.5), PptInches(1.55))
            title_tf = title_box.text_frame
            title_tf.word_wrap = True
            title_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tp = title_tf.paragraphs[0]
            tp.text = str(item.get("title") or title)
            _set_ppt_paragraph_font(tp, 30, True, navy)

            subtitle_box = slide.shapes.add_textbox(PptInches(0.98), PptInches(4.15), PptInches(8.8), PptInches(0.6))
            sp = subtitle_box.text_frame.paragraphs[0]
            sp.text = str(item.get("subtitle") or "科研与工程智能分析")
            _set_ppt_paragraph_font(sp, 17, False, muted)

            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.98), PptInches(5.15), PptInches(2.1), PptInches(0.06))
            rule.fill.solid()
            rule.fill.fore_color.rgb = cyan
            rule.line.fill.background()
        else:
            topbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PptInches(0.12))
            topbar.fill.solid()
            topbar.fill.fore_color.rgb = blue
            topbar.line.fill.background()

            number = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(0.72), PptInches(0.58), PptInches(0.58), PptInches(0.38))
            number.fill.solid()
            number.fill.fore_color.rgb = navy
            number.line.fill.background()
            np = number.text_frame.paragraphs[0]
            np.alignment = PP_ALIGN.CENTER
            np.text = f"{index:02d}"
            _set_ppt_paragraph_font(np, 11, True, white)

            title_box = slide.shapes.add_textbox(PptInches(1.5), PptInches(0.47), PptInches(10.8), PptInches(0.68))
            title_tf = title_box.text_frame
            title_tf.word_wrap = True
            tp = title_tf.paragraphs[0]
            tp.text = str(item.get("title") or f"第 {index} 部分")
            _set_ppt_paragraph_font(tp, 25, True, navy)

            divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, PptInches(0.72), PptInches(1.35), PptInches(11.85), PptInches(0.015))
            divider.fill.solid()
            divider.fill.fore_color.rgb = line
            divider.line.fill.background()

            body_box = slide.shapes.add_textbox(PptInches(0.95), PptInches(1.72), PptInches(11.35), PptInches(4.75))
            body = body_box.text_frame
            body.clear()
            body.word_wrap = True
            body.margin_left = PptInches(0.08)
            body.margin_right = PptInches(0.08)
            bullets = item.get("bullets") or []
            if isinstance(bullets, str):
                bullets = [bullets]
            for bullet_index, bullet in enumerate(bullets[:8]):
                paragraph = body.paragraphs[0] if bullet_index == 0 else body.add_paragraph()
                paragraph.text = f"•  {str(bullet)}"
                paragraph.level = 0
                paragraph.space_after = PptPt(13)
                paragraph.line_spacing = 1.14
                _set_ppt_paragraph_font(paragraph, 18, False, ink)

            callout = item.get("takeaway")
            if callout:
                callout_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PptInches(8.6), PptInches(5.82), PptInches(3.7), PptInches(0.65))
                callout_box.fill.solid()
                callout_box.fill.fore_color.rgb = white
                callout_box.line.color.rgb = line
                cp = callout_box.text_frame.paragraphs[0]
                cp.text = str(callout)
                _set_ppt_paragraph_font(cp, 10, False, muted)

        footer = slide.shapes.add_textbox(PptInches(0.72), PptInches(7.05), PptInches(11.9), PptInches(0.22))
        fp = footer.text_frame.paragraphs[0]
        fp.text = f"EnerTri V4.2    {index + 1:02d} / {len(slide_data):02d}"
        fp.alignment = PP_ALIGN.RIGHT
        _set_ppt_paragraph_font(fp, 8.5, False, muted)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path

